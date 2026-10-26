
"""Core functions for CareerNexus AI."""

import os
import re


SKILL_ALIASES = {
    "python": ["python"],
    "machine learning": ["machine learning", "ml"],
    "deep learning": ["deep learning", "dl"],
    "generative ai": ["generative ai", "genai", "generative artificial intelligence"],
    "llms": ["llm", "llms", "large language model", "large language models"],
    "rag": ["rag", "retrieval augmented generation", "retrieval-augmented generation"],
    "fastapi": ["fastapi", "fast api"],
    "flask": ["flask"],
    "rest apis": ["rest api", "rest apis", "restful api", "restful apis"],
    "sql": ["sql"],
    "mysql": ["mysql"],
    "sqlite": ["sqlite"],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch", "torch"],
    "git": ["git"],
    "github": ["github"],
    "docker": ["docker"],
    "postman": ["postman"],
    "jwt authentication": ["jwt", "json web token", "json web tokens"],
    "object-oriented programming": [
        "object-oriented programming", "object oriented programming", "oop"
    ],
    "data analysis": ["data analysis", "data analytics"],
    "data visualization": ["data visualization", "data visualisation"],
    "feature engineering": ["feature engineering"],
    "model evaluation": ["model evaluation", "model validation"],
    "api testing": ["api testing"],
    "react": ["react", "react.js", "reactjs"],
    "javascript": ["javascript"],
    "java": ["java"],
    "c++": ["c++"],
    "natural language processing": ["natural language processing", "nlp"],
    "transformers": ["transformers", "hugging face transformers"],
    "faiss": ["faiss"],
    "chroma": ["chroma", "chromadb"],
    "langchain": ["langchain"],
    "streamlit": ["streamlit"],
    "pytest": ["pytest"],
    "unit testing": ["unit testing", "unit tests"],
    "data structures and algorithms": [
        "data structures and algorithms", "data structures", "algorithms", "dsa"
    ],
    "html": ["html"],
    "css": ["css"],
    "linux": ["linux"],
    "aws": ["aws", "amazon web services"],
    "google cloud": ["google cloud", "gcp"],
}


def _normalise(text):
    return re.sub(r"\s+", " ", str(text or "").lower()).strip()


def _contains_term(text, term):
    pattern = (
        r"(?<![a-z0-9+#])"
        + re.escape(term.lower())
        + r"(?![a-z0-9+#])"
    )
    return re.search(pattern, text.lower()) is not None


def extract_skills(text):
    normalised = _normalise(text)
    found = []

    for skill, aliases in SKILL_ALIASES.items():
        if any(_contains_term(normalised, alias) for alias in aliases):
            found.append(skill)

    return sorted(set(found))


def match_skills(resume_text, job_description):
    resume_skills = extract_skills(resume_text)
    required_skills = extract_skills(job_description)

    resume_set = set(resume_skills)
    required_set = set(required_skills)

    matched = sorted(resume_set & required_set)
    missing = sorted(required_set - resume_set)

    coverage = (
        round(len(matched) / len(required_set) * 100)
        if required_set else 0
    )

    return {
        "match_percentage": coverage,
        "coverage_percentage": coverage,
        "matched_skills": matched,
        "missing_skills": missing,
        "resume_skills": resume_skills,
        "required_skills": required_skills,
    }


def analyze_match(resume_text, job_description):
    result = match_skills(resume_text, job_description)
    missing = result["missing_skills"]
    matched = result["matched_skills"]

    recommendations = []

    if missing:
        recommendations.append(
            "Review these job requirements and add evidence to your resume "
            "only if you genuinely have the skills: " + ", ".join(missing) + "."
        )

    if matched:
        recommendations.append(
            "Strengthen your project and experience bullets with evidence for: "
            + ", ".join(matched) + "."
        )

    if not result["required_skills"]:
        recommendations.append(
            "No recognised skills were found. Check that the complete job "
            "description was provided."
        )

    recommendations.append(
        "Do not claim skills or experience you cannot demonstrate."
    )

    result["recommendations"] = recommendations
    result["skills_to_review"] = missing
    result["note"] = (
        "This is a keyword-based skills coverage estimate, not a validated "
        "ATS score or hiring prediction."
    )
    return result


def extract_text(uploaded_file):
    filename = getattr(uploaded_file, "name", "").lower()
    raw_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        import fitz

        parts = []
        with fitz.open(stream=raw_bytes, filetype="pdf") as pdf:
            for page in pdf:
                parts.append(page.get_text())

        extracted = "\n".join(parts)

    elif filename.endswith(".docx"):
        import io
        from docx import Document

        doc = Document(io.BytesIO(raw_bytes))
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        extracted = "\n".join(parts)

    elif filename.endswith((".txt", ".md", ".csv")):
        extracted = raw_bytes.decode("utf-8", errors="replace")

    else:
        raise ValueError("Upload a PDF, DOCX, TXT, Markdown, or CSV file.")

    if not extracted.strip():
        raise ValueError(
            "No readable text was extracted. Scanned PDFs may need OCR."
        )

    return extracted.strip()


def chunk_text(text, chunk_size=1200, overlap=200):
    """Split text into overlapping chunks without discarding document text."""
    text = str(text or "").strip()

    if not text:
        return []

    chunk_size = max(200, int(chunk_size))
    overlap = max(0, min(int(overlap), chunk_size - 1))

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - overlap

    return chunks


def build_documents(resume_text, job_description, chunk_size=1200, overlap=200):
    """Create source-labelled chunks from the full resume and job description."""
    documents = []

    for source, source_text in (
        ("Resume", resume_text),
        ("Job Description", job_description),
    ):
        for index, chunk in enumerate(
            chunk_text(source_text, chunk_size, overlap)
        ):
            documents.append({
                "text": chunk,
                "source": source,
                "chunk_id": index,
                "metadata": {
                    "source": source,
                    "chunk_id": index,
                },
            })

    return documents


def load_embedding_model(
    model_name="sentence-transformers/all-MiniLM-L6-v2"
):
    """Load embeddings when available; keyword retrieval remains the fallback."""
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer(model_name)
    except (ImportError, OSError, RuntimeError):
        return None


def _question_terms(text):
    stop_words = {
        "the", "and", "for", "with", "what", "which", "when", "where",
        "how", "does", "this", "that", "are", "you", "your", "from",
        "have", "has", "can", "could", "would", "should", "about",
        "tell", "give", "please", "based", "resume", "job",
    }

    return {
        word for word in re.findall(r"[a-zA-Z0-9+#.-]+", _normalise(text))
        if len(word) > 1 and word not in stop_words
    }


def _keyword_score(question, text):
    question_words = _question_terms(question)
    text_words = set(
        re.findall(r"[a-zA-Z0-9+#.-]+", _normalise(text))
    )

    if not question_words:
        return 0.0

    overlap = len(question_words & text_words)
    return overlap / len(question_words)


def retrieve_chunks(
    question,
    documents,
    top_k=8,
    embedding_model=None,
    per_source_minimum=2,
):
    """
    Retrieve relevant chunks, preserving resume and job-description coverage.

    Embedding similarity is preferred when available. Keyword scoring is used
    when embeddings are unavailable or fail. A small number of chunks from
    each source is prioritised to reduce one-source-only results.
    """
    if not documents:
        return []

    top_k = max(1, int(top_k))
    per_source_minimum = max(0, int(per_source_minimum))

    if embedding_model is not None:
        try:
            import numpy as np

            texts = [str(doc.get("text", "")) for doc in documents]
            query_vector = embedding_model.encode(
                [question], normalize_embeddings=True
            )[0]
            vectors = embedding_model.encode(
                texts, normalize_embeddings=True
            )
            scores = np.dot(vectors, query_vector).tolist()
        except Exception:
            scores = [
                _keyword_score(question, doc.get("text", ""))
                for doc in documents
            ]
    else:
        scores = [
            _keyword_score(question, doc.get("text", ""))
            for doc in documents
        ]

    ranked_indices = sorted(
        range(len(documents)),
        key=lambda index: scores[index],
        reverse=True,
    )

    selected = []
    selected_indices = set()
    source_counts = {}

    # First, include relevant chunks from each source when available.
    source_names = []
    for doc in documents:
        source = doc.get("source", "Document")
        if source not in source_names:
            source_names.append(source)

    for source in source_names:
        source_indices = [
            index for index in ranked_indices
            if documents[index].get("source", "Document") == source
        ]

        for index in source_indices[:per_source_minimum]:
            if len(selected) >= top_k:
                break
            if index not in selected_indices:
                selected.append(index)
                selected_indices.add(index)
                source_counts[source] = source_counts.get(source, 0) + 1

    # Fill remaining places with the highest-ranked unused chunks.
    for index in ranked_indices:
        if len(selected) >= top_k:
            break
        if index not in selected_indices:
            selected.append(index)
            selected_indices.add(index)

    return [
        {
            **documents[index],
            "relevance_score": round(float(scores[index]), 4),
        }
        for index in selected
    ]


def load_generation_model(model_name=None):
    """Create the Gemini client configuration."""
    from google import genai

    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        try:
            import streamlit as st
            api_key = st.secrets.get("GEMINI_API_KEY")
        except Exception:
            api_key = None

    selected_model = (
        model_name
        or os.getenv("GEMINI_MODEL")
        or "gemini-3.5-flash-lite"
    )

    if not api_key:
        return {"client": None, "model_name": selected_model}

    return {
        "client": genai.Client(api_key=api_key),
        "model_name": selected_model,
    }


def _generate_with_gemini(prompt, generation_model=None):
    config = generation_model or load_generation_model()
    client = config.get("client")
    model_name = config.get("model_name", "gemini-3.5-flash-lite")

    if client is None:
        return None

    try:
        response = client.models.generate_content(
            model=model_name,
            contents=prompt,
        )
        return (response.text or "").strip() or None
    except Exception:
        return None


def answer_question(
    question,
    documents=None,
    top_k=8,
    embedding_model=None,
    generator=None,
    generation_model=None,
):
    """Answer a career question using retrieved resume and job context."""
    question = str(question or "").strip()

    if not question:
        return {
            "answer": "Please enter a question.",
            "sources": [],
            "retrieved_chunks": [],
        }

    retrieved = retrieve_chunks(
        question=question,
        documents=documents or [],
        top_k=top_k,
        embedding_model=embedding_model,
    )

    sources = []
    context_parts = []

    for item in retrieved:
        source = item.get("source", "Document")
        chunk = item.get("text", "")

        if source not in sources:
            sources.append(source)

        context_parts.append(
            f"[Source: {source}; section chunk {item.get('chunk_id', '?')}]\n"
            f"{chunk}"
        )

    context = "\n\n".join(context_parts)

    if not context:
        context = "No resume or job-description context was supplied."

    prompt = f"""
You are CareerNexus AI, a career assistant.

Use the provided context to answer the user's question.
The context can contain multiple excerpts from a resume and a job description.
Do not claim that the resume is incomplete unless the provided excerpts clearly
show that it is incomplete. Retrieval excerpts may cover only part of a document.
Do not assume a skill is missing merely because it does not appear in a retrieved
excerpt. Distinguish between:
- explicitly supported by the context,
- not found in the retrieved context,
- genuinely absent from the full resume (only if the full resume is available).

Give practical, concise advice. Never invent qualifications or experience.
When asked about skills, compare the resume evidence with the job requirements.
When the available context is insufficient, state the limitation plainly.

RETRIEVED CONTEXT:
{context}

USER QUESTION:
{question}
"""

    answer = None

    if callable(generator):
        try:
            answer = generator(prompt)
        except Exception:
            answer = None

    if not answer:
        answer = _generate_with_gemini(prompt, generation_model)

    if not answer:
        answer = (
            "I couldn't generate a Gemini response. Check GEMINI_API_KEY and "
            "the Gemini model configuration. The relevant text retrieved was:\n\n"
            + "\n\n".join(context_parts[:3])
        )

    return {
        "answer": str(answer),
        "sources": sources,
        "retrieved_chunks": retrieved,
    }
